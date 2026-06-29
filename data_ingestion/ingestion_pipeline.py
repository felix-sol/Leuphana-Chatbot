import os
from pathlib import Path
from typing import Optional


class IngestionPipeline:
    def __init__(self, data_dir: str = "./raw_input_data", chunk_size: int = 1000, overlap: int = 100):
        self.data_dir = Path(data_dir)
        self.chunk_size = chunk_size
        self.overlap = overlap

    # ── Parsers ──────────────────────────────────────────────────────────────

    def parse_docx(self, path: Path) -> str:
        from docx import Document
        doc = Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

    def parse_pdf(self, path: Path) -> str:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        return "\n".join(
            page.extract_text() or "" for page in reader.pages
        )

    def parse_pptx(self, path: Path) -> str:
        from pptx import Presentation
        prs = Presentation(str(path))
        texts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text_frame.text)
        return "\n".join(t for t in texts if t.strip())

    def _parse_file(self, path: Path) -> Optional[str]:
        suffix = path.suffix.lower()
        try:
            if suffix == ".docx":
                return self.parse_docx(path)
            elif suffix == ".pdf":
                return self.parse_pdf(path)
            elif suffix == ".pptx":
                return self.parse_pptx(path)
        except Exception as e:
            print(f"[WARNING] Konnte {path.name} nicht lesen: {e}")
        return None

    # ── Document loading ─────────────────────────────────────────────────────

    def _get_category(self, path: Path) -> str:
        """Leitet die Kategorie aus der Ordnerstruktur ab."""
        try:
            relative = path.relative_to(self.data_dir)
            return relative.parts[0] if relative.parts else "Allgemein"
        except ValueError:
            return "Allgemein"

    def load_all_documents(self) -> list[dict]:
        """
        Liest alle .docx, .pdf und .pptx Dateien aus data_dir.
        Gibt eine Liste von Dicts mit 'text', 'source' und 'category' zurück.
        """
        supported = {".docx", ".pdf", ".pptx"}
        documents = []

        for path in sorted(self.data_dir.rglob("*")):
            if path.suffix.lower() not in supported:
                continue
            text = self._parse_file(path)
            if not text or not text.strip():
                continue
            documents.append({
                "text": text,
                "source": path.name,
                "category": self._get_category(path),
                "full_path": str(path),
            })

        print(f"[INFO] {len(documents)} Dokumente geladen.")
        return documents

    # ── Chunking ─────────────────────────────────────────────────────────────

    def chunk_text(self, text: str) -> list[str]:
        """Teilt Text in überlappende Chunks auf (zeichenbasiert, max ~512 Tokens)."""
        chunks = []
        start = 0
        while start < len(text):
            end = start + self.chunk_size
            chunk = text[start:end].strip()
            if chunk:
                chunks.append(chunk)
            start += self.chunk_size - self.overlap
        return chunks

    def chunk_documents(self, documents: list[dict]) -> list[dict]:
        """
        Chunked alle Dokumente und gibt eine flache Liste von Chunks mit Metadaten zurück.
        Jeder Chunk: {'text', 'source', 'category', 'chunk_index'}
        """
        all_chunks = []
        for doc in documents:
            chunks = self.chunk_text(doc["text"])
            for i, chunk in enumerate(chunks):
                all_chunks.append({
                    "text": chunk,
                    "source": doc["source"],
                    "category": doc["category"],
                    "chunk_index": i,
                })
        print(f"[INFO] {len(all_chunks)} Chunks erstellt.")
        return all_chunks


if __name__ == "__main__":
    pipeline = IngestionPipeline()
    docs = pipeline.load_all_documents()
    chunks = pipeline.chunk_documents(docs)
    print(f"Beispiel-Chunk:\n{chunks[0]['text'][:200]}")